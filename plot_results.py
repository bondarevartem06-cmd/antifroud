"""
Визуализация результатов compare_models.py (файл results_summary.csv или любой <prefix>_summary.csv).

Примеры (запуск из терминала PyCharm, одной строкой):
  python plot_results.py
  python plot_results.py --file run2_summary.csv --out_dir plots_run2
  python plot_results.py --plot_set with_fp --models lgbm xgb --scenarios plain cat_w
  python plot_results.py --show            # дополнительно открыть окна с графиками

  python plot_results.py --table_only      # только таблица со всей информацией (без графиков)

Таблица (table_all.csv и table_all.xlsx) - все метрики по всем сценариям, наборам и моделям,
отсортированные по чистой выгоде, с разницей относительно базового сценария и базового набора.
В Excel-файле ячейки окрашены: зелёный = лучше, красный = хуже. Здесь (и только здесь) сравниваются
наборы признаков, в том числе без FP-Growth и с ним.

Графики сохраняются в папку --out_dir (по умолчанию plots/) в виде PNG:
Графики строятся для ОДНОГО набора признаков (--plot_set; по умолчанию последний в файле):
  <метрика>_by_scenario.png    - влияние ВЕСОВ: столбцы = сценарии весов, по оси X модели
  fit_sec_by_scenario.png, time_vs_quality.png - время обучения и его связь с качеством
  overfit_gap.png, precision_vs_recall.png, net_benefit_vs_cap.png - переобучение, размен precision/recall, бюджет проверок
  net_benefit_*.png, cost_total_*.png - экономика (стоимость проверок, пропусков, ложных блокировок)
  cost_breakdown.png           - из чего состоят потери
  tradeoff_money_vs_false_alarms.png - сэкономленные деньги против ложных срабатываний
"""
import argparse
import os
import sys

import matplotlib

if "--show" not in sys.argv:  # без --show рисуем в файлы, окна не открываем
    matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.lines import Line2D
from matplotlib.ticker import FuncFormatter

METRICS = {  # колонка: (заголовок, формат подписей над столбцами)
    "PR-AUC": ("PR-AUC (больше = лучше)", "{:.3f}"),
    "F1": ("F1 при лимите службы мониторинга", "{:.3f}"),
    "detected": ("Обнаружено мошеннических транзакций", "{:,.0f}"),
    "saved_money": ("Сэкономленные деньги (сумма обнаруженных мошеннических платежей)", "{:,.0f}"),
    "net_benefit": ("Чистая выгода относительно «ничего не проверять», у.е. (больше = лучше)", "{:,.0f}"),
    "cost_total": ("Суммарные потери: проверки + пропуски + ложные блокировки, у.е. (меньше = лучше)", "{:,.0f}"),
    "fit_sec": ("Время обучения модели, секунды (меньше = лучше)", "{:,.1f}"),
    "false_alarms": ("Ложные срабатывания (честные транзакции помечены как мошенничество)", "{:,.0f}"),
}
MARKERS = ["o", "s", "^", "D", "v", "P"]


def thousands(ax):
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.0f}"))


def bars_by_facet(df, metric, facet, hue, x, path, labels=True):
    """Сгруппированные столбцы: по оси X - x, цвет - hue, отдельная панель на каждое значение facet."""
    title, fmt = METRICS[metric]
    panels = list(df[facet].unique())
    fig, axes = plt.subplots(1, len(panels), figsize=(max(9.5, 5.8 * len(panels) + 1), 5), squeeze=False, sharey=True)
    for ax, p in zip(axes[0], panels):
        pv = df[df[facet] == p].pivot_table(index=x, columns=hue, values=metric, aggfunc="mean")
        pv = pv.reindex(index=list(df[x].unique()), columns=[c for c in df[hue].unique() if c in pv.columns])
        pv.plot(kind="bar", ax=ax, width=0.8, rot=0, legend=False)
        ax.set_title(f"{facet}: {p}", fontsize=10)
        ax.set_xlabel("модель")
        ax.grid(axis="y", alpha=0.3)
        if metric not in ("PR-AUC", "F1", "fit_sec"):
            thousands(ax)
        if labels:
            for c in ax.containers:
                ax.bar_label(c, labels=[fmt.format(v) for v in c.datavalues], fontsize=7, padding=2, rotation=90)
            ax.margins(y=0.18)
    axes[0][0].set_ylabel(metric)
    h, l = axes[0][0].get_legend_handles_labels()
    fig.legend(h, l, title=hue, loc="upper right")
    fig.suptitle(title, fontsize=12)
    fig.tight_layout(rect=(0, 0, 0.93, 0.95))
    fig.savefig(path, dpi=150)


def tradeoff(df, path):
    """Точка = (сценарий, модель): правее = больше ложных срабатываний, выше = больше сэкономлено."""
    panels = list(df["set"].unique())
    scen, models = list(df["scenario"].unique()), list(df["model"].unique())
    colors = {s: plt.cm.tab10(i % 10) for i, s in enumerate(scen)}
    marks = {m: MARKERS[i % len(MARKERS)] for i, m in enumerate(models)}
    fig, axes = plt.subplots(1, len(panels), figsize=(max(9.5, 5.8 * len(panels) + 1), 5), squeeze=False, sharey=True)
    for ax, p in zip(axes[0], panels):
        for _, r in df[df["set"] == p].iterrows():
            ax.scatter(r["false_alarms"], r["saved_money"], s=90, color=colors[r["scenario"]],
                       marker=marks[r["model"]], edgecolor="black", linewidth=0.5)
        ax.set_title(f"set: {p}", fontsize=10)
        ax.set_xlabel("ложные срабатывания (шт.)")
        ax.grid(alpha=0.3)
        thousands(ax)
        ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.0f}"))
    axes[0][0].set_ylabel("сэкономленные деньги")
    handles = [Line2D([], [], marker="o", color="w", markerfacecolor=colors[s], markeredgecolor="black",
                      markersize=9, label=s) for s in scen]
    handles += [Line2D([], [], marker=marks[m], color="w", markerfacecolor="lightgray", markeredgecolor="black",
                       markersize=9, label=m) for m in models]
    fig.legend(handles=handles, title="сценарий / модель", loc="upper right")
    fig.suptitle("Сэкономленные деньги против ложных срабатываний (левее и выше = лучше)", fontsize=12)
    fig.tight_layout(rect=(0, 0, 0.9, 0.95))
    fig.savefig(path, dpi=150)


def cost_breakdown(df, path):
    """Из чего состоят потери: проверки, пропущенное мошенничество, ложные блокировки (пунктир - без проверки)."""
    parts = [("cost_review", "проверки аналитиком"), ("cost_missed", "пропущенное мошенничество"),
             ("cost_false_block", "ложные блокировки")]
    panels = list(df["set"].unique())
    widths = [max(6.0, 0.6 * (df["set"] == p).sum()) for p in panels]
    fig, axes = plt.subplots(1, len(panels), figsize=(max(10, sum(widths) + 1.5), 5.8), squeeze=False, sharey=True)
    for ax, p in zip(axes[0], panels):
        d = df[df["set"] == p].reset_index(drop=True)
        bottom = pd.Series(0.0, index=d.index)
        for col, label in parts:
            ax.bar(d.index, d[col], bottom=bottom, label=label, width=0.75)
            bottom += d[col]
        ax.axhline(d["cost_no_model"].mean(), color="black", linestyle="--", linewidth=1,
                   label="потери без проверки вообще")
        ax.set_xticks(d.index)
        ax.set_xticklabels(d["scenario"] + "\n" + d["model"], fontsize=7, rotation=90)
        ax.set_title(f"set: {p}", fontsize=10)
        ax.grid(axis="y", alpha=0.3)
        thousands(ax)
    axes[0][0].set_ylabel("у.е.")
    h, l = axes[0][0].get_legend_handles_labels()
    fig.legend(h, l, loc="lower center", ncol=4, fontsize=8)
    fig.suptitle("Структура потерь (ниже = лучше; ниже пунктира = проверка окупается)", fontsize=12)
    fig.tight_layout(rect=(0, 0.07, 1, 0.95))
    fig.savefig(path, dpi=150)


TABLE_COLUMNS = [  # порядок столбцов в полной таблице
    "scenario", "set", "model", "cat_weights", "amt_weights",
    "PR-AUC", "ROC-AUC", "precision", "recall", "F1",
    "flagged", "flagged_share", "frauds_total", "detected", "missed", "false_alarms",
    "cards_fraud_total", "cards_detected",
    "saved_money", "missed_money", "saved_share", "blocked_legit_money",
    "cost_review", "cost_missed", "cost_false_block", "cost_total", "cost_no_model", "net_benefit", "fit_sec", "sec",
]
KEY_COLUMNS = ["rank_net_benefit", "scenario", "set", "model", "PR-AUC", "precision", "detected",
               "false_alarms", "saved_money", "cost_total", "net_benefit", "fit_sec"]
DELTA_METRICS = ["PR-AUC", "detected", "saved_money", "net_benefit"]


def build_table(df):
    """Полная таблица: все метрики + место по выгоде + разница с базовым сценарием и базовым набором."""
    t = df.copy()
    first_sc, first_set = t["scenario"].iloc[0], t["set"].iloc[0]  # базовые: первые в файле (обычно plain)
    t = t[[c for c in TABLE_COLUMNS if c in t.columns] + [c for c in t.columns if c not in TABLE_COLUMNS]]
    for m in [m for m in DELTA_METRICS if m in t.columns]:
        b_sc = t[t["scenario"] == first_sc].groupby(["set", "model"])[m].mean()
        b_set = t[t["set"] == first_set].groupby(["scenario", "model"])[m].mean()
        t[f"d_{m}_vs_scn_{first_sc}"] = t[m].to_numpy() - b_sc.reindex(
            pd.MultiIndex.from_frame(t[["set", "model"]])).to_numpy()
        t[f"d_{m}_vs_set_{first_set}"] = t[m].to_numpy() - b_set.reindex(
            pd.MultiIndex.from_frame(t[["scenario", "model"]])).to_numpy()
    sort_col = "net_benefit" if "net_benefit" in t.columns else "PR-AUC"
    t.insert(0, f"rank_{sort_col}".replace("-", "_"), t[sort_col].rank(ascending=False, method="min").astype(int))
    t = t.sort_values(t.columns[0], kind="stable").reset_index(drop=True)
    money = [c for c in t.columns if any(k in c for k in ("money", "cost", "benefit"))]
    t[money] = t[money].round(2)
    other = [c for c in t.select_dtypes("float").columns if c not in money]
    t[other] = t[other].round(4)
    return t


HIGHER_BETTER = {"PR-AUC", "ROC-AUC", "precision", "recall", "F1", "detected", "cards_detected",
                 "saved_money", "saved_share", "net_benefit"}  # чем больше, тем лучше
LOWER_BETTER = {"missed", "false_alarms", "missed_money", "blocked_legit_money", "cost_review", "cost_missed",
                "cost_false_block", "cost_total", "fit_sec", "sec", "gap_fit_val", "gap_val_test"}  # чем меньше, тем лучше
RED, YELLOW, GREEN = "F8696B", "FFEB84", "63BE7B"


def paint(ws, data):
    """Цветовая шкала по каждому столбцу: чем лучше значение, тем зеленее; чем хуже, тем краснее.
    Столбцы-названия, веса и постоянные столбцы (например, число помеченных операций) не красятся."""
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter
    for i, name in enumerate(data.columns, start=1):
        col = pd.to_numeric(data[name], errors="coerce")
        if col.notna().sum() < 2 or col.nunique() < 2:  # нечего красить
            continue
        if name.startswith("d_") or name in HIGHER_BETTER:
            lo, hi = RED, GREEN
        elif name.startswith("rank_") or name in LOWER_BETTER:
            lo, hi = GREEN, RED
        else:
            continue
        letter = get_column_letter(i)
        rng = f"{letter}2:{letter}{ws.max_row}"
        if name.startswith("d_") and col.min() < 0 < col.max():  # разницы: ноль = нейтральный белый
            rule = ColorScaleRule(start_type="min", start_color=lo, mid_type="num", mid_value=0,
                                  mid_color="FFFFFF", end_type="max", end_color=hi)
        else:
            rule = ColorScaleRule(start_type="min", start_color=lo, mid_type="percentile", mid_value=50,
                                  mid_color=YELLOW, end_type="max", end_color=hi)
        ws.conditional_formatting.add(rng, rule)


def save_table(t, out_dir):
    """Сохраняет table_all.csv и (если установлен openpyxl) table_all.xlsx с листами all и key.
    В Excel-файле ячейки окрашены: зелёный = лучше, красный = хуже (CSV цветов не хранит)."""
    paths = [os.path.join(out_dir, "table_all.csv")]
    t.to_csv(paths[0], index=False, encoding="utf-8-sig")
    try:
        from openpyxl.utils import get_column_letter
        import openpyxl  # noqa: F401
    except ImportError:
        print("(Excel-файл не создан: pip install openpyxl; CSV сохранён)")
        return paths
    xlsx = os.path.join(out_dir, "table_all.xlsx")
    key = [c for c in KEY_COLUMNS if c in t.columns] + [c for c in t.columns if c.startswith(("d_net_benefit", "d_PR_AUC", "d_PR-AUC"))]
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        for name, data in (("all", t), ("key", t[key])):  # all = все столбцы, key = главные
            data.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = f"{get_column_letter(list(data.columns).index('model') + 2)}2"  # шапка и столбцы до model
            ws.auto_filter.ref = ws.dimensions
            for cells in ws.columns:
                w = max(len(str(c.value)) if c.value is not None else 0 for c in cells[:60])
                ws.column_dimensions[cells[0].column_letter].width = min(max(10, w + 2), 40)
            paint(ws, data)
    paths.append(xlsx)
    return paths


def overfit_chart(df, path):
    """PR-AUC на обучающей части, валидации и test: большой разрыв fit -> val означает переобучение."""
    panels = list(df["set"].unique())
    widths = [max(6.0, 0.6 * (df["set"] == p).sum()) for p in panels]
    fig, axes = plt.subplots(1, len(panels), figsize=(max(10, sum(widths) + 1.5), 5.8), squeeze=False, sharey=True)
    for ax, p in zip(axes[0], panels):
        d = df[df["set"] == p].reset_index(drop=True)
        x = d.index.to_numpy()
        for k, (col, label) in enumerate((("PR-AUC_fit", "обучение (fit)"), ("PR-AUC_val", "валидация"),
                                          ("PR-AUC", "test"))):
            ax.bar(x + (k - 1) * 0.27, d[col], width=0.27, label=label)
        ax.set_xticks(x)
        ax.set_xticklabels(d["scenario"] + "\n" + d["model"], fontsize=7, rotation=90)
        ax.set_title(f"set: {p}", fontsize=10)
        ax.grid(axis="y", alpha=0.3)
    axes[0][0].set_ylabel("PR-AUC")
    h, l = axes[0][0].get_legend_handles_labels()
    fig.legend(h, l, loc="lower center", ncol=3, fontsize=8)
    fig.suptitle("Контроль переобучения: PR-AUC на fit / val / test (большой разрыв fit-val = переобучение)", fontsize=11)
    fig.tight_layout(rect=(0, 0.07, 1, 0.95))
    fig.savefig(path, dpi=150)


def curve_charts(cv, out_dir, mark=0.003):
    """precision против recall и выгода при разных бюджетах проверок (cap = доля операций в сутки)."""
    panels = list(cv["set"].unique())
    scen, models = list(cv["scenario"].unique()), list(cv["model"].unique())
    colors = {s_: plt.cm.tab10(i % 10) for i, s_ in enumerate(scen)}
    marks = {m: MARKERS[i % len(MARKERS)] for i, m in enumerate(models)}
    pct = FuncFormatter(lambda v, _: f"{v * 100:.2g}%")
    jobs = (("recall", "precision", "Precision против recall при разных бюджетах проверок",
             "precision_vs_recall.png"),
            ("cap", "net_benefit", "Чистая выгода при разных бюджетах проверок (пунктир = текущий лимит)",
             "net_benefit_vs_cap.png"))
    saved = []
    for xcol, ycol, title, name in jobs:
        if ycol not in cv.columns:
            continue
        fig, axes = plt.subplots(1, len(panels), figsize=(max(9.5, 5.8 * len(panels) + 1.5), 5.2), squeeze=False, sharey=True)
        for ax, p in zip(axes[0], panels):
            sub = cv[cv["set"] == p]
            best = sub.groupby(["scenario", "model"])["recall"].max().idxmax()  # лучшая линия: на ней подписи лимитов
            for (sc, m), d in sub.groupby(["scenario", "model"], sort=False):
                d = d.sort_values("cap")
                ax.plot(d[xcol], d[ycol], marker=marks[m], color=colors[sc], linewidth=1.2, markersize=5)
                if xcol == "recall" and (sc, m) == best:
                    for _, r in d.iterrows():
                        ax.annotate(f"{r['cap'] * 100:.2g}%", (r["recall"], r["precision"]), fontsize=7,
                                    textcoords="offset points", xytext=(5, 5))
            ax.set_title(f"set: {p}", fontsize=10)
            ax.grid(alpha=0.3)
            if xcol == "cap":
                ax.set_xscale("log")
                ax.set_xticks(sorted(cv["cap"].unique()))
                ax.xaxis.set_major_formatter(pct)
                ax.xaxis.set_minor_formatter(FuncFormatter(lambda v, _: ""))
                ax.axvline(mark, color="black", linestyle="--", linewidth=1)
                ax.axhline(0, color="gray", linewidth=0.8)
                ax.set_xlabel("бюджет проверок (доля операций в сутки)")
                thousands(ax)
            else:
                ax.set_xlabel("recall (доля найденного мошенничества); подписи у точек = бюджет проверок лучшей линии")
        axes[0][0].set_ylabel("precision" if ycol == "precision" else "чистая выгода, у.е.")
        handles = [Line2D([], [], marker="o", color="w", markerfacecolor=colors[x], markeredgecolor="black",
                          markersize=9, label=x) for x in scen]
        handles += [Line2D([], [], marker=marks[m], color="w", markerfacecolor="lightgray", markeredgecolor="black",
                           markersize=9, label=m) for m in models]
        fig.legend(handles=handles, title="сценарий / модель", loc="upper right")
        fig.suptitle(title, fontsize=11)
        fig.tight_layout(rect=(0, 0, 0.9, 0.95))
        path = os.path.join(out_dir, name)
        fig.savefig(path, dpi=150)
        saved.append(path)
    return saved


def time_vs_quality(df, path):
    """Время обучения против качества: левее и выше = быстрее и лучше."""
    ycol = "net_benefit" if "net_benefit" in df.columns else "PR-AUC"
    scen, models = list(df["scenario"].unique()), list(df["model"].unique())
    colors = {s_: plt.cm.tab10(i % 10) for i, s_ in enumerate(scen)}
    marks = {m: MARKERS[i % len(MARKERS)] for i, m in enumerate(models)}
    fig, ax = plt.subplots(figsize=(9, 5.5))
    for _, r in df.iterrows():
        ax.scatter(max(r["fit_sec"], 1e-3), r[ycol], s=90, color=colors[r["scenario"]],
                   marker=marks[r["model"]], edgecolor="black", linewidth=0.5)
    ax.set_xscale("log")
    ax.set_xlabel("время обучения, секунды (логарифмическая шкала)")
    ax.set_ylabel("чистая выгода, у.е." if ycol == "net_benefit" else "PR-AUC")
    ax.grid(alpha=0.3)
    if ycol == "net_benefit":
        thousands(ax)
    handles = [Line2D([], [], marker="o", color="w", markerfacecolor=colors[s_], markeredgecolor="black",
                      markersize=9, label=s_) for s_ in scen]
    handles += [Line2D([], [], marker=marks[m], color="w", markerfacecolor="lightgray", markeredgecolor="black",
                       markersize=9, label=m) for m in models]
    fig.legend(handles=handles, title="сценарий / модель", loc="upper right")
    fig.suptitle("Время обучения и качество (левее и выше = быстрее и лучше)", fontsize=12)
    fig.tight_layout(rect=(0, 0, 0.86, 0.95))
    fig.savefig(path, dpi=150)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--file", default="results_summary.csv")
    ap.add_argument("--out_dir", default="plots")
    ap.add_argument("--sets", nargs="+", help="оставить только эти наборы признаков (и в таблице, и на графиках)")
    ap.add_argument("--models", nargs="+", help="показать только эти модели")
    ap.add_argument("--scenarios", nargs="+", help="показать только эти сценарии весов")
    ap.add_argument("--plot_set", help="набор признаков для графиков (по умолчанию последний в файле); "
                                       "в таблице всегда все наборы")
    ap.add_argument("--no_labels", action="store_true", help="не писать числа над столбцами")
    ap.add_argument("--show", action="store_true", help="открыть окна с графиками")
    ap.add_argument("--curve_file", help="файл кривых <prefix>_curve.csv (по умолчанию выводится из --file)")
    ap.add_argument("--no_table", action="store_true", help="не создавать таблицу table_all")
    ap.add_argument("--table_only", action="store_true", help="только таблица, без графиков")
    ap.add_argument("--top", type=int, default=15, help="сколько строк таблицы печатать в консоль")
    args = ap.parse_args()

    if not os.path.exists(args.file):
        raise SystemExit(f"Нет файла {args.file}. Сначала запустите compare_models.py "
                         f"(он создаёт results_summary.csv).")
    df = pd.read_csv(args.file)
    if "scenario" not in df.columns:  # старый формат results.csv без сценариев
        df["scenario"] = "plain"
    for col, wanted in (("set", args.sets), ("model", args.models), ("scenario", args.scenarios)):
        if wanted:
            df = df[df[col].isin(wanted)]
            if df.empty:
                raise SystemExit(f"После фильтра --{col}s ничего не осталось. "
                                 f"Проверьте названия (регистр важен).")

    df_all = df  # таблица: все наборы признаков (в том числе без FP и с FP)
    plot_set = args.plot_set or df["set"].iloc[-1]  # графики: ОДИН набор признаков
    if plot_set not in set(df["set"]):
        raise SystemExit(f"Нет набора '{plot_set}'. В файле есть: {list(df['set'].unique())}")
    df = df[df["set"] == plot_set]

    os.makedirs(args.out_dir, exist_ok=True)
    if not args.no_table:
        table = build_table(df_all)
        show = [c for c in KEY_COLUMNS if c in table.columns] + [
            c for c in table.columns if c.startswith(("d_net_benefit", "d_PR-AUC"))]
        print(f"Лучшие {min(args.top, len(table))} из {len(table)} строк (сортировка по выгоде):")
        with pd.option_context("display.width", 250, "display.max_columns", None):
            print(table[show].head(args.top).to_string(index=False))
        print("Таблица со всей информацией:", ", ".join(save_table(table, args.out_dir)))
        if args.table_only:
            return
    print(f"Графики построены для набора '{plot_set}' (в таблице все наборы: {', '.join(df_all['set'].unique())})")

    saved = []
    metrics = [m for m in METRICS if m in df.columns]
    for m in metrics:  # влияние весов (панель одна: выбранный набор)
        path = os.path.join(args.out_dir, f"{m}_by_scenario.png")
        bars_by_facet(df, m, facet="set", hue="scenario", x="model", path=path, labels=not args.no_labels)
        saved.append(path)
    if {"saved_money", "false_alarms"} <= set(df.columns):
        path = os.path.join(args.out_dir, "tradeoff_money_vs_false_alarms.png")
        tradeoff(df, path)
        saved.append(path)
    if {"cost_review", "cost_missed", "cost_false_block", "cost_no_model"} <= set(df.columns):
        path = os.path.join(args.out_dir, "cost_breakdown.png")
        cost_breakdown(df, path)
        saved.append(path)
    if "fit_sec" in df.columns:  # время обучения против качества
        path = os.path.join(args.out_dir, "time_vs_quality.png")
        time_vs_quality(df, path)
        saved.append(path)
    if "PR-AUC_fit" in df.columns:  # контроль переобучения
        path = os.path.join(args.out_dir, "overfit_gap.png")
        overfit_chart(df, path)
        saved.append(path)
    curve_file = args.curve_file or args.file.replace("_summary.csv", "_curve.csv")
    if os.path.exists(curve_file):  # кривые precision / recall / выгода по бюджету проверок
        cv = pd.read_csv(curve_file)
        cv = cv[cv["set"] == plot_set]
        for col, wanted in (("model", args.models), ("scenario", args.scenarios)):
            if wanted:
                cv = cv[cv[col].isin(wanted)]
        if not cv.empty:
            saved += curve_charts(cv, args.out_dir)

    print("Сохранено графиков:", len(saved))
    for p in saved:
        print("  ", p)
    if args.show:
        plt.show()
    else:
        plt.close("all")


if __name__ == "__main__":
    main()
